#!/usr/bin/env python3
"""Build clean review DOCX files from the frozen Markdown manuscript sources."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 55, start: int = 65, bottom: int = 55, end: int = 65) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))


def configure_document(doc: Document, *, supplementary: bool) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(9.8 if supplementary else 10.5)
    normal.paragraph_format.line_spacing = 1.07 if supplementary else 1.08
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, before, after, color in (
        ("Title", 16, 0, 10, "17365D"),
        ("Heading 1", 13, 10, 4, "17365D"),
        ("Heading 2", 11.5, 8, 3, "1F4E79"),
        ("Heading 3", 10.5, 6, 2, "2F5597"),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(9.8 if supplementary else 10.5)
        style.paragraph_format.space_after = Pt(2)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("MINI ARTICHOKES — SUPPLEMENT" if supplementary else "MINI ARTICHOKES")
    run.font.name = "Arial"
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(89, 89, 89)
    add_page_number(section.footer.paragraphs[0])


def clean_inline(text: str) -> str:
    text = text.replace("\\quad", " ").replace("\\Longleftrightarrow", "⇔")
    text = text.replace("\\ne", "≠").replace("\\ge", "≥").replace("\\le", "≤")
    text = text.replace("\\sum", "Σ").replace("\\Pr", "Pr").replace("\\mathrm", "")
    text = text.replace("\\text", "").replace("\\sim", "~")
    text = text.replace("\\{", "{").replace("\\}", "}")
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    return text


def add_inline(paragraph, text: str) -> None:
    text = clean_inline(text)
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(64, 64, 64)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def is_separator(row: list[str]) -> bool:
    return bool(row) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row)


def add_table(doc: Document, markdown_lines: list[str]) -> None:
    rows = [row for row in parse_table(markdown_lines) if not is_separator(row)]
    if not rows:
        return
    # Keep a table's caption/lead-in with its header, and keep the header with
    # the first data row. This avoids orphaned headers at a page boundary.
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    font_size = 7.3 if width >= 7 else 7.8 if width >= 5 else 8.4
    for i, row in enumerate(rows):
        table_row = table.rows[i]
        tr_pr = table_row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if i == 0:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            tr_pr.append(repeat_header)
        for j in range(width):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            # Short manuscript tables should move as a unit instead of leaving
            # unlabeled trailing rows at the top of the next page. Longer
            # supplementary tables may paginate, with their header repeated.
            if len(rows) <= 8 and i < len(rows) - 1:
                paragraph.paragraph_format.keep_with_next = True
            if i == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                set_cell_shading(cell, "D9EAF7")
            add_inline(paragraph, row[j] if j < len(row) else "")
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(font_size)
                if i == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.0
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    paragraph._p.get_or_add_pPr().append(shading)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(7.8)


def render_markdown(source: Path, destination: Path, *, supplementary: bool) -> dict[str, int]:
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc, supplementary=supplementary)
    i = 0
    tables = 0
    code_blocks = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
            code_blocks += 1
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            tables += 1
            continue
        if stripped in {"\\[", "$$"}:
            close = "\\]" if stripped == "\\[" else "$$"
            block = []
            i += 1
            while i < len(lines) and lines[i].strip() != close:
                block.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            run = p.add_run(clean_inline(" ".join(block)))
            run.font.name = "Cambria Math"
            run.italic = True
            i += 1
            continue
        if stripped.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, stripped[2:])
            i += 1
            continue
        heading = re.match(r"^(##|###|####)\s+(.+)$", stripped)
        if heading:
            level = {"##": 1, "###": 2, "####": 3}[heading.group(1)]
            doc.add_heading(clean_inline(heading.group(2)), level=level)
            i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            add_inline(p, stripped)
            i += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            add_inline(p, stripped[2:])
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(64, 64, 64)
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("#", "|", "```", "- ", "> ", "\\[", "$$")):
                break
            if re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    check = Document(destination)
    text = "\n".join(p.text for p in check.paragraphs)
    required_values = (
        ("58", "51", ".01953125", "withheld-by-control", "Python34 five-pair", "C++26 five-pair", "Preservation-contract checklist", "Go39", "33:31", "+20.51", "JavaScript49", "47", "HumanEvalFix", "164", "3.13.11", "test-covered denominator", "132/178", "S21", "[0,0,-1,-1,0]", "S23", "S24", "QuixBugs")
        if supplementary
        else ("58/90", "51/90", "p=.01953", "withheld-by-control", "210/300", "190/300", "59/90", "54/90", "Go39", "33/39", "+20.51", "JavaScript49", "47/49", "HumanEvalFix", "164/164", "3.13.11", "32/38", "132/178", "8:1", "does not repeat", "QuixBugs")
    )
    for required in required_values:
        if required not in text and not any(required in cell.text for table in check.tables for row in table.rows for cell in row.cells):
            raise RuntimeError(f"missing required content after DOCX build: {required}")
    return {"paragraphs": len(check.paragraphs), "tables": len(check.tables), "code_blocks": code_blocks}


def main() -> int:
    outputs = [
        (
            ROOT / "Mini_Artichokes_manuscript_v4.md",
            ROOT / "output" / "Mini_Artichokes_Main_Revised_v25.docx",
            False,
        ),
        (
            ROOT / "Mini_Artichokes_supplementary_v2.md",
            ROOT / "output" / "Mini_Artichokes_Supplementary_Revised_v25.docx",
            True,
        ),
    ]
    for source, destination, supplementary in outputs:
        result = render_markdown(source, destination, supplementary=supplementary)
        print(destination)
        print(result)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
